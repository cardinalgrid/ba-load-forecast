"""Monthly installed solar PV capacity behind each BA's load, from two public EIA sources.

    python scripts/pv_capacity.py

* Net-metered PV (mostly rooftop, behind the meter): Form EIA-861M, ``Monthly_Totals-States`` sheet, total PV
  capacity (MW AC) by state and month, 2016 onward; the states of a BA are summed (ISO-NE: CT, MA, ME, NH, RI,
  VT; CAISO: CA; FPL: FL).
* Utility-scale PV in front of the meter: Form EIA-860M generator inventory (operating and retired sheets),
  solar photovoltaic generators by balancing authority code, nameplate MW cumulated by operating month and
  reduced by retirements.

Writes ``docs/pv_capacity_monthly.csv`` (ba, month, net_metered_mw, utility_mw, total_mw). Months before the
first EIA-861M month carry its value; the files are cached in ``data/pv/``. Capacity is a slow, publicly known
quantity, so a forecaster at the 9 a.m. cut-off knows it; the models use it through ``sky.pv_estimate``.
"""

from __future__ import annotations

import sys
import urllib.request
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "data" / "pv"
OUT = ROOT / "docs" / "pv_capacity_monthly.csv"
BA_STATES = {"ISNE": ("CT", "MA", "ME", "NH", "RI", "VT"), "CISO": ("CA",), "FPL": ("FL",)}
NET_METERING = {
    2016: "archive/xls/f826netmetering2016.xlsx", 2017: "archive/xls/net_metering2017.xlsx",
    2018: "archive/xls/net_metering2018.xlsx", 2019: "archive/xls/net_metering2019.xlsx",
    2020: "archive/xls/net_metering2020.xlsx", 2021: "archive/xls/net_metering_2021.xlsx",
    2022: "archive/xls/net_metering2022.xlsx", 2023: "archive/xls/net_metering2023.xlsx",
    2024: "archive/xls/net_metering_2024.xlsx", 2025: "archive/xls/net_metering2025.xlsx",
    2026: "xls/net_metering2026.xlsx",
}
GENERATORS = "https://www.eia.gov/electricity/data/eia860m/xls/july_generator2026.xlsx"


def fetch(url: str, name: str) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    f = CACHE / name
    if not f.exists():
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        f.write_bytes(urllib.request.urlopen(req, timeout=300).read())
    return f


def net_metered(year: int) -> pd.DataFrame:
    f = fetch(f"https://www.eia.gov/electricity/data/eia861m/{NET_METERING[year]}", f"net_metering{year}.xlsx")
    ws = openpyxl.load_workbook(f, read_only=True, data_only=True)["Monthly_Totals-States"]
    rows = [r for r in ws.iter_rows(values_only=True)]
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Year")
    d = pd.DataFrame(rows[header + 1:], columns=range(len(rows[header])))
    d = d.rename(columns={0: "year", 1: "month", 2: "state", 8: "pv_total_mw"})[["year", "month", "state", "pv_total_mw"]]
    d = d[d["year"].astype(str).str.fullmatch(r"\d{4}")]
    d["pv_total_mw"] = pd.to_numeric(d["pv_total_mw"], errors="coerce")
    d["month"] = pd.to_datetime(d["year"].astype(int).astype(str) + "-" + d["month"].astype(int).astype(str).str.zfill(2) + "-01")
    return d[["month", "state", "pv_total_mw"]]


def utility_scale() -> pd.DataFrame:
    f = fetch(GENERATORS, "generator2026.xlsx")
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    parts = []
    for sheet in ("Operating", "Retired"):
        rows = list(wb[sheet].iter_rows(values_only=True))
        header = next(i for i, r in enumerate(rows) if r and r[0] == "Entity ID")
        d = pd.DataFrame(rows[header + 1:], columns=rows[header])
        d = d[d["Technology"] == "Solar Photovoltaic"]
        d["ba"] = d["Balancing Authority Code"].astype(str).str.strip()
        d["mw"] = pd.to_numeric(d["Nameplate Capacity (MW)"], errors="coerce")
        d["start"] = pd.to_datetime(d["Operating Year"].astype(str) + "-" + d["Operating Month"].astype(str) + "-01", errors="coerce")
        if sheet == "Retired":
            d["end"] = pd.to_datetime(d["Retirement Year"].astype(str) + "-" + d["Retirement Month"].astype(str) + "-01", errors="coerce")
        else:
            d["end"] = pd.NaT
        parts.append(d[["ba", "mw", "start", "end"]])
    return pd.concat(parts).dropna(subset=["mw", "start"])


def main() -> int:
    nm = pd.concat([net_metered(y) for y in NET_METERING])
    months = pd.date_range("2015-01-01", nm["month"].max(), freq="MS")
    gens = utility_scale()
    out = []
    for ba, states in BA_STATES.items():
        s = nm[nm["state"].isin(states)].groupby("month")["pv_total_mw"].sum(min_count=1).reindex(months).ffill().bfill()
        g = gens[gens["ba"] == ba]
        u = pd.Series([(g[(g["start"] <= m) & (g["end"].isna() | (g["end"] > m))]["mw"]).sum() for m in months], index=months)
        out.append(pd.DataFrame({"ba": ba, "month": months, "net_metered_mw": s.round(1).to_numpy(), "utility_mw": u.round(1).to_numpy()}))
    t = pd.concat(out, ignore_index=True)
    t["total_mw"] = (t["net_metered_mw"] + t["utility_mw"]).round(1)
    t.to_csv(OUT, index=False, date_format="%Y-%m")
    print(t[t["month"].dt.month == 7].pivot(index="month", columns="ba", values="total_mw").to_string())
    print(f"written {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
